import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from threading import Lock

from fastapi import FastAPI, HTTPException
from openpyxl import Workbook, load_workbook
import uvicorn


WORKBOOK_PATH = Path(__file__).with_name("reservations.xlsx")
HEADERS = (
    "reservation_id",
    "user_id",
    "passenger_name",
    "membership_level",
    "trip_type",
    "origin",
    "destination",
    "flight_number",
    "departure_time",
    "cabin_class",
    "payment_method",
    "created_time",
    "travel_insurance",
    "airline_cancelled",
    "already_flown",
    "reservation_status",
)
_lock = Lock()
app = FastAPI(title="Reservation API")


def seed_reservations() -> None:
    """Write the demo workbook. Re-running this replaces the file."""
    now = datetime.now(timezone.utc).replace(microsecond=0)
    rows = [
        [
            "3FRNFB",
            "amelia_sanchez_4739",
            "Amelia Sanchez",
            "regular",
            "one way",
            "MCO",
            "CLT",
            "AA482",
            (now + timedelta(days=12)).isoformat(),
            "basic economy",
            "credit_card",
            (now - timedelta(days=30)).isoformat(),
            "no",
            "no",
            "no",
            "active",
        ],
        [
            "8JK2LQ",
            "jordan_lee_1001",
            "Jordan Lee",
            "silver",
            "one way",
            "BOS",
            "DFW",
            "DL210",
            (now + timedelta(days=6)).isoformat(),
            "economy",
            "credit_card",
            (now - timedelta(hours=2)).isoformat(),
            "no",
            "no",
            "no",
            "active",
        ],
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "reservations"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(row)
    workbook.save(WORKBOOK_PATH)


def _rows() -> list[dict[str, str]]:
    if not WORKBOOK_PATH.is_file():
        raise HTTPException(
            status_code=500,
            detail="reservations.xlsx is missing. Run: python reservation_server.py --seed",
        )
    workbook = load_workbook(WORKBOOK_PATH)
    sheet = workbook["reservations"]
    records = []
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        if cells[0] is None:
            continue
        records.append(
            {header: "" if value is None else str(value) for header, value in zip(HEADERS, cells)}
        )
    return records


def _write(records: list[dict[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "reservations"
    sheet.append(list(HEADERS))
    for record in records:
        sheet.append([record.get(header, "") for header in HEADERS])
    workbook.save(WORKBOOK_PATH)


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/reservations")
def get_reservation_details(user_id: str) -> list[dict[str, str]]:
    with _lock:
        return [row for row in _rows() if row["user_id"] == user_id]


@app.post("/reservations/{reservation_id}/cancel")
def cancel_reservation(reservation_id: str) -> dict[str, str]:
    """Cancel without checking policy. The agent must decide before calling this."""
    with _lock:
        records = _rows()
        for record in records:
            if record["reservation_id"] != reservation_id:
                continue
            if record["reservation_status"] == "cancelled":
                raise HTTPException(status_code=409, detail="Reservation is already cancelled")
            record["reservation_status"] = "cancelled"
            _write(records)
            return record
    raise HTTPException(status_code=404, detail="Reservation not found")


if __name__ == "__main__":
    if "--seed" in sys.argv:
        seed_reservations()
        print(f"Wrote {WORKBOOK_PATH}")
    if "--seed" in sys.argv and "--serve" not in sys.argv:
        raise SystemExit(0)
    uvicorn.run(app, host="127.0.0.1", port=8000)
